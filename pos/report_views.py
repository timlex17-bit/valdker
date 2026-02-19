import io
import json
import openpyxl

from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404
from django.template.loader import render_to_string, get_template
from django.db.models import Sum, F
from django.db.models.functions import TruncMonth, TruncDate
from django.utils import timezone

from xhtml2pdf import pisa

from .models import Order, OrderItem, Expense, Shop


def expense_chart_view(request):
    data = (
        Expense.objects
        .annotate(month=TruncMonth("date"))
        .values("month")
        .annotate(total=Sum("amount"))
        .order_by("month")
    )

    labels = [d["month"].strftime("%B %Y") for d in data if d.get("month")]
    totals = [float(d["total"] or 0) for d in data if d.get("month")]

    # kalau template kamu Chart.js, ini lebih aman:
    return render(request, "pos/expense_chart.html", {
        "labels": json.dumps(labels),
        "totals": json.dumps(totals),
    })


def expense_report_view(request):
    expenses = Expense.objects.all().order_by("-date", "-time", "-id")
    return render(request, "pos/expense_report.html", {"expenses": expenses})


def sales_chart_view(request):
    data = (
        Order.objects
        .annotate(month=TruncMonth("created_at"))
        .values("month")
        .annotate(
            total_sales=Sum("total"),
            total_discount=Sum("discount"),
            total_tax=Sum("tax"),
        )
        .order_by("month")
    )

    labels = [d["month"].strftime("%b %Y") for d in data if d.get("month")]
    sales = [float(d["total_sales"] or 0) for d in data if d.get("month")]
    discounts = [float(d["total_discount"] or 0) for d in data if d.get("month")]
    taxes = [float(d["total_tax"] or 0) for d in data if d.get("month")]

    total_order_price = sum(sales)
    total_discount = sum(discounts)
    total_tax = sum(taxes)
    net_sales = total_order_price - total_discount + total_tax

    return render(request, "pos/sales_chart.html", {
        "labels": json.dumps(labels),
        "sales": json.dumps(sales),
        "total_order_price": total_order_price,
        "total_discount": total_discount,
        "total_tax": total_tax,
        "net_sales": net_sales,
    })


def sales_report_view(request):
    month = request.GET.get("month")
    rows = []

    orders = Order.objects.prefetch_related("items__product", "items__weight_unit", "customer").order_by("-created_at")

    if month:
        try:
            year, month_number = map(int, month.split("-"))
            orders = orders.filter(created_at__year=year, created_at__month=month_number)
        except ValueError:
            pass

    for order in orders:
        for item in order.items.all():
            rows.append({
                "product_name": item.product.name if item.product else "-",
                "invoice_id": f"INV{order.id:015d}",
                "qty": item.quantity,
                "weight": item.weight_unit.name if item.weight_unit else "-",
                "total_price": (item.quantity or 0) * (item.price or 0),
                "order_date": timezone.localtime(order.created_at).strftime("%d %B, %Y"),
            })

    return render(request, "pos/sales_report.html", {"rows": rows})


def sales_export_view(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    ws.append([
        "Transaction ID", "Date", "Customer", "Product",
        "Qty", "Price", "Subtotal", "Discount", "Tax", "Total"
    ])

    orders = Order.objects.prefetch_related("items__product", "customer").order_by("-created_at")

    for order in orders:
        for item in order.items.all():
            ws.append([
                f"INV{order.id:015d}",
                timezone.localtime(order.created_at).strftime("%Y-%m-%d %H:%M"),
                order.customer.name if order.customer else "-",
                item.product.name if item.product else "-",
                item.quantity,
                item.price,
                (item.quantity or 0) * (item.price or 0),
                order.discount,
                order.tax,
                order.total
            ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=sales_report.xlsx"
    wb.save(response)
    return response


def order_pdf_view(request, order_id):
    order = get_object_or_404(
        Order.objects.prefetch_related("items__product", "items__weight_unit").select_related("served_by", "customer"),
        id=order_id
    )
    shop = Shop.objects.first()

    # bikin receipt_items biar cocok dengan template receipt yang kamu pakai di views.py
    receipt_items = []
    for item in order.items.all():
        qty = item.quantity or 0
        unit_price = item.price or 0
        receipt_items.append({
            "name": item.product.name if item.product else "-",
            "qty": qty,
            "unit_price": unit_price,
            "line_total": qty * unit_price,
            "order_type": getattr(item, "order_type", ""),
            "order_type_icon": "",
        })

    template = get_template("pos/order_receipt.html")
    html = template.render({"order": order, "shop": shop, "receipt_items": receipt_items})

    result = io.BytesIO()
    pdf = pisa.pisaDocument(io.BytesIO(html.encode("UTF-8")), result)

    if not pdf.err:
        return HttpResponse(result.getvalue(), content_type="application/pdf")
    return HttpResponse("PDF generation error", status=500)


def sales_report_pdf_view(request):
    month = request.GET.get("month")
    orders = Order.objects.prefetch_related("items__product", "customer").order_by("-created_at")

    if month:
        year, m = map(int, month.split("-"))
        orders = orders.filter(created_at__year=year, created_at__month=m)

    html = render_to_string("pos/sales_report_pdf.html", {"orders": orders})
    result = io.BytesIO()
    pisa.CreatePDF(io.BytesIO(html.encode("UTF-8")), dest=result)
    return HttpResponse(result.getvalue(), content_type="application/pdf")


def sales_report_excel_view(request):
    month = request.GET.get("month")
    orders = Order.objects.prefetch_related("items__product", "items__weight_unit", "customer").order_by("-created_at")

    if month:
        year, m = map(int, month.split("-"))
        orders = orders.filter(created_at__year=year, created_at__month=m)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"
    ws.append(["Product Name", "Invoice ID", "Qty", "Weight", "Total Price", "Order Date"])

    for order in orders:
        for item in order.items.all():
            ws.append([
                item.product.name if item.product else "-",
                f"INV{order.id:015d}",
                item.quantity,
                item.weight_unit.name if item.weight_unit else "",
                (item.quantity or 0) * (item.price or 0),
                timezone.localtime(order.created_at).strftime("%Y-%m-%d %H:%M"),
            ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=sales_report.xlsx"
    wb.save(response)
    return response


def sales_report_print_view(request):
    month = request.GET.get("month")
    orders = Order.objects.prefetch_related("items__product", "items__weight_unit").order_by("-created_at")

    if month:
        year, m = map(int, month.split("-"))
        orders = orders.filter(created_at__year=year, created_at__month=m)

    rows = []
    for order in orders:
        for item in order.items.all():
            rows.append({
                "product_name": item.product.name if item.product else "-",
                "invoice_id": f"INV{order.id:015d}",
                "qty": item.quantity,
                "weight": item.weight_unit.name if item.weight_unit else "",
                "total_price": (item.quantity or 0) * (item.price or 0),
                "order_date": timezone.localtime(order.created_at).date(),
            })

    return render(request, "pos/sales_report_print.html", {"rows": rows})


# --- OPTIONAL dashboards (kalau kamu memang mau tetap pakai dari report_urls.py) ---
def daily_profit_dashboard_view(request):
    # ini boleh tetap panggil yang di views.py kalau kamu mau,
    # tapi lebih aman kalau kamu pindahkan logic dashboard ke sini.
    return render(request, "pos/daily_profit_dashboard.html")


def monthly_pl_dashboard_view(request):
    return render(request, "pos/monthly_pl_dashboard.html")
